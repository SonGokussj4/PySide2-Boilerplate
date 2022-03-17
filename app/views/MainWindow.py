from pathlib import Path
from PySide2.QtWidgets import QMainWindow, QActionGroup, QDesktopWidget, QMessageBox, QApplication, QAction, QMenu
from PySide2.QtGui import Qt, QCloseEvent, QResizeEvent
from PySide2.QtCore import QSettings, QPoint, QSize, QTranslator, QCoreApplication, QEvent, QLocale, Slot

import openpyxl as xlsx

from .SettingsWindow import SettingsWindow
from ..ui.MainWindow_ui import Ui_MainWindow
from ..helpers.utils import get_language_from_code, load_translations
from ..helpers.logging import setup_logger
from ..helpers import constants, win_functions


logger = setup_logger(__name__)
logger.debug('This MainWindow message should appear on the console')


class MainWindow(QMainWindow):
    """Main Window."""

    def __init__(self, parent=None):
        super().__init__(parent)

        self.ui = Ui_MainWindow()
        self.ui.setupUi(self)

        # Settings
        self.settings = QSettings(constants.APP_NAME, "General")
        logger.debug(f"Settings filepath: {self.settings.fileName()}")

        # Windows size
        win_functions.setSizeAndPosition(self)

        # Fix QAction gruping
        action_group = QActionGroup(self)
        action_group.addAction(self.ui.action_English)
        action_group.addAction(self.ui.action_Czech)

        # Slots and Signals
        # self.ui.btnOpenSettings.clicked.connect(self.openSettingsWindow)

        # Actions
        self.ui.action_English.triggered.connect(lambda: self.changeLanguage('en_US'))
        self.ui.action_Czech.triggered.connect(lambda: self.changeLanguage('cs_CZ'))

        # Set language
        self.translator = QTranslator()

        # Default values
        self.UIComponents()

    def UIComponents(self):
        """Set UI components."""
        default_language_code = QLocale.system().name()
        language_code = self.settings.value("language", default_language_code, type=str)
        self.changeLanguage(language_code)

        # Set checkbox of selected language
        language_name = get_language_from_code(language_code)
        language_object = [self.ui.action_English, self.ui.action_Czech][["English", "Czech"].index(language_name)]
        language_object.setChecked(True)

        # Initial values
        self.loadInitialValues()

    def changeLanguage(self, language_code):
        logger.debug(f"Changing language to: {language_code}")
        self.settings.setValue("language", language_code)

        # Set language
        load_translations(translator=self.translator, code=language_code)
        self.ui.retranslateUi(self)

    @Slot(str)
    def on_cboMatGroup_currentTextChanged(self, text):
        print("on_cboMatGroup_currentTextChanged")
        self._updateCboMatSub(text)

    def loadInitialValues(self):
        """Load initial values."""
        # TEST PURPOSES
        self.ui.cboMatGroup.addItems(("A", "B", "C"))

        # open xlsx file
        workbook = xlsx.load_workbook(Path("DATA/example1.xlsx"))
        print(f'workbook: {workbook}')
        # get all sheets
        sheets = workbook.sheetnames
        print(f'sheets: {sheets}')

        csv_ws = workbook['csv']
        csv_tables = csv_ws.tables
        print([tbl for tbl in csv_tables])

        structures_ws = workbook['Structures']
        structures_tables = structures_ws.tables
        # https://openpyxl.readthedocs.io/en/latest/worksheet_tables.html
        first_table = structures_tables["tblMAT_STD"]
        print(f'first_table: {first_table}')
        print(f'first_table.name: {first_table.name}')
        columns = [col.name for col in first_table.tableColumns]
        print(f'columns[{type(columns)}]: {columns}')
        data = structures_ws[first_table.ref]
        content = [[cell.value for cell in row] for row in data]
        header = content[0]
        rest = content[1:]
        print(f'header: {header}')
        print(f'rest: {rest}')

    def _updateCboMatSub(self, text):
        self.ui.cboMatSub.clear()
        if text == "A":
            self.ui.cboMatSub.addItems(("1", "2", "3"))
        elif text == "B":
            self.ui.cboMatSub.addItems(("4", "5", "6"))
        elif text == "C":
            self.ui.cboMatSub.addItems(("7", "8", "9"))

    @Slot()
    def on_BtnOpenSettings_clicked(self):
        self.settingsWindow = SettingsWindow()
        self.settingsWindow.ui.txtSetText.setText(self.ui.lblSimpleText.text())
        self.settingsWindow.show()

    @Slot()
    def on_actionPreferences_triggered(self):
        """Open settings window."""
        print("Action Preferences triggered")
        self.openSettingsWindow()

    # Close window on ESC key
    def keyPressEvent(self, event):
        if event.key() == Qt.Key_Escape:
            self.close()

    def changeEvent(self, event):
        if event.type() == QEvent.LanguageChange:
            self.ui.retranslateUi(self)
        super(MainWindow, self).changeEvent(event)

    def closeEvent(self, event: QCloseEvent) -> None:
        """Close window."""
        # close = QMessageBox()
        # close.setText(QApplication.translate('MainWindow', 'You Sure?'))
        # close.setStandardButtons(QMessageBox.Yes | QMessageBox.Cancel)
        # close = close.exec()

        # if close != QMessageBox.Yes:
        #     event.ignore()
        #     return

        event.accept()

        win_name = self.objectName()
        self.settings.setValue(f"{win_name}/size", self.size())
        self.settings.setValue(f"{win_name}/pos", self.pos())

        return super().closeEvent(event)
